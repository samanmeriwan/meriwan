#!/usr/bin/env python3
"""Build the Mariwan master-data import files (tasks 4.4 / 15.1).

Usage:
    python3 tools/build_import_files.py [--items "Categories and item details.xlsx"] [--out docs/data_load]
                                        [--cost-currency AED]

Writes blank Odoo import templates (with an instructions sheet, allowed-value
dropdowns and duplicate / missing-value highlighting) and, when the client item
list is given, ready-to-import product and price list files built from it.
All files use the column names of the standard Odoo import (Settings > Import).
"""
import argparse
import os

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RED = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
HEADER = Font(bold=True)

CATEGORIES = ['Brake Pads', 'Brake Rotors', 'Drum Brake Shoe', 'Brake Lines', 'Brake Fluid',
              'Callipers', 'Brake Wire Sensor', 'Big Brake Upgrade Kit']

# template name -> list of (odoo column, description, mandatory, allowed values or None, example)
TEMPLATES = {
    'products': [
        ('id', 'Unique external ID, e.g. product_16_1184', True, None, 'product_16_1184'),
        ('default_code', 'Internal reference (SKU) - must be unique', True, None, '16-1184'),
        ('name', 'Product name', True, None, 'Power Stop Z16 Evolution Brake Pads'),
        ('categ_id', 'Product category', True, CATEGORIES, 'Brake Pads'),
        ('type', 'Product type', True, ['Goods', 'Service', 'Combo'], 'Goods'),
        ('is_storable', 'Track inventory', True, ['TRUE', 'FALSE'], 'TRUE'),
        ('barcode', 'Barcode (optional, unique)', False, None, ''),
        ('uom_id', 'Unit of measure', True, ['Units', 'Pieces', 'Carton', 'Pallet', 'Container'], 'Units'),
        ('mariwan_brand', 'Brand', False, None, 'Power Stop'),
        ('mariwan_model', 'Model', False, None, 'Z16'),
        ('hs_code', 'HS code (customs)', False, None, '8708.30'),
        ('country_of_origin', 'Country of origin', False, None, 'United States'),
        ('weight', 'Weight (kg) - used for landed cost by weight', False, None, '2.5'),
        ('volume', 'Volume (m3) - used for landed cost by volume', False, None, '0.01'),
        ('tracking', 'Tracking (leave empty for none)', False, ['lot', 'serial'], ''),
        ('mariwan_fulfilment_type', 'Default fulfilment', True, ['stock', 'import', 'dropship'], 'stock'),
        ('standard_price', 'Cost (AED)', False, None, '45'),
        ('list_price', 'Sales price (AED) - UAE price', False, None, '54'),
        ('mariwan_min_stock', 'Minimum stock (alert only)', False, None, '5'),
        ('seller_ids/partner_id', 'Main vendor name (must exist as a vendor)', False, None, 'Powerstop'),
        ('seller_ids/price', 'Vendor price', False, None, '45'),
        ('seller_ids/currency_id', 'Vendor price currency', False, ['AED', 'USD', 'EUR', 'CNY'], 'AED'),
        ('seller_ids/delay', 'Vendor lead time (days)', False, None, '30'),
    ],
    'customers': [
        ('id', 'Unique external ID, e.g. customer_0001', True, None, 'customer_0001'),
        ('name', 'Customer / company name', True, None, 'Al Noor Auto Spare Parts LLC'),
        ('is_company', 'Company', True, ['TRUE', 'FALSE'], 'TRUE'),
        ('customer_rank', 'Set 1 for customers', True, ['1'], '1'),
        ('mariwan_customer_type', 'Customer type', True, ['cash', 'credit'], 'credit'),
        ('mariwan_region', 'Sales region', True,
         ['abu_dhabi', 'dubai', 'sharjah', 'ajman', 'umm_al_quwain', 'ras_al_khaimah', 'fujairah', 'gcc', 'other'],
         'dubai'),
        ('property_product_pricelist', 'Price list', True, ['UAE Market', 'GCC Market'], 'UAE Market'),
        ('credit_limit', 'Credit limit (AED) - credit customers', False, None, '50000'),
        ('property_payment_term_id', 'Payment terms', True,
         ['Immediate Payment', '7 Days', '15 Days', '30 Days', '45 Days', '60 Days',
          'Advance 50% + Balance on Delivery'], '30 Days'),
        ('street', 'Street', False, None, 'Al Quoz Industrial Area 3'),
        ('city', 'City', False, None, 'Dubai'),
        ('country_id', 'Country', True, None, 'United Arab Emirates'),
        ('phone', 'Phone / WhatsApp', False, None, '+971 4 000 0000'),
        ('email', 'Email', False, None, 'info@example.ae'),
        ('mariwan_trade_licence', 'Trade licence number', False, None, '123456'),
    ],
    'vendors': [
        ('id', 'Unique external ID, e.g. vendor_powerstop', True, None, 'vendor_powerstop'),
        ('name', 'Vendor name', True, None, 'Powerstop'),
        ('is_company', 'Company', True, ['TRUE', 'FALSE'], 'TRUE'),
        ('supplier_rank', 'Set 1 for vendors', True, ['1'], '1'),
        ('mariwan_vendor_type', 'Vendor type', True, ['overseas', 'local', 'dropship'], 'overseas'),
        ('country_id', 'Country', True, None, 'United States'),
        ('property_purchase_currency_id', 'Purchase currency', True, ['AED', 'USD', 'EUR', 'CNY'], 'USD'),
        ('property_supplier_payment_term_id', 'Payment terms', False, None, '30 Days'),
        ('email', 'Email', False, None, 'orders@example.com'),
        ('phone', 'Phone', False, None, '+1 000 000 0000'),
        ('bank_ids/account_number', 'Bank account number / IBAN', False, None, ''),
    ],
    'vendor_prices': [
        ('partner_id', 'Vendor name', True, None, 'Powerstop'),
        ('product_tmpl_id', 'Product (internal reference or name)', True, None, '16-1184'),
        ('product_code', 'Vendor product code', False, None, '16-1184'),
        ('min_qty', 'Minimum quantity', False, None, '1'),
        ('price', 'Unit price', True, None, '45'),
        ('currency_id', 'Currency', True, ['AED', 'USD', 'EUR', 'CNY'], 'AED'),
        ('delay', 'Lead time (days)', False, None, '30'),
        ('date_start', 'Valid from (YYYY-MM-DD)', False, None, ''),
        ('date_end', 'Valid to (YYYY-MM-DD)', False, None, ''),
    ],
    'opening_stock': [
        ('product_id', 'Product (internal reference)', True, None, '16-1184'),
        ('location_id', 'Location', True, ['MRW/Stock'], 'MRW/Stock'),
        ('lot_id', 'Lot / serial (tracked products only)', False, None, ''),
        ('inventory_quantity', 'Counted quantity', True, None, '20'),
        ('standard_price', 'Unit cost (AED) - set on the product before applying', False, None, '45'),
    ],
    'opening_balances': [
        ('account_id', 'Account code', True, None, '101401'),
        ('partner_id', 'Partner (receivable / payable lines)', False, None, ''),
        ('name', 'Label', True, None, 'Opening balance'),
        ('debit', 'Debit (AED)', True, None, '100000'),
        ('credit', 'Credit (AED)', True, None, '0'),
    ],
}

UNIQUE_COLUMNS = {'id', 'default_code', 'barcode'}


def _write_template(path, columns, rows=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Import'
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.font = HEADER
    if rows is None:
        ws.append([c[4] for c in columns])
    else:
        for row in rows:
            ws.append(row)
    last = max(ws.max_row, 1000)
    for idx, (name, _desc, mandatory, allowed, _ex) in enumerate(columns, start=1):
        letter = openpyxl.utils.get_column_letter(idx)
        ws.column_dimensions[letter].width = max(14, len(name) + 4)
        rng = f'{letter}2:{letter}{last}'
        if allowed:
            dv = DataValidation(type='list', formula1='"%s"' % ','.join(allowed), allow_blank=not mandatory)
            dv.error = f'Choose one of: {", ".join(allowed)}'
            ws.add_data_validation(dv)
            dv.add(rng)
        if mandatory:
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'AND(COUNTA($A2:${openpyxl.utils.get_column_letter(len(columns))}2)>0,ISBLANK({letter}2))'],
                fill=RED))
        if name in UNIQUE_COLUMNS:
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'AND({letter}2<>"",COUNTIF(${letter}$2:${letter}${last},{letter}2)>1)'], fill=RED))
    info = wb.create_sheet('Instructions')
    info.append(['Column', 'Description', 'Mandatory', 'Allowed values', 'Example'])
    for cell in info[1]:
        cell.font = HEADER
    for name, desc, mandatory, allowed, example in columns:
        info.append([name, desc, 'Yes' if mandatory else 'No', ', '.join(allowed or []), example])
    info.append([])
    info.append(['Red cells = missing mandatory value or duplicate. Fix them before importing.'])
    info.append(['Import in Odoo: open the list view > gear icon > Import records > upload the "Import" sheet.'])
    for col, width in zip('ABCDE', (28, 55, 11, 45, 30)):
        info.column_dimensions[col].width = width
    wb.save(path)


def _xmlid(sku):
    return 'product_' + ''.join(ch if ch.isalnum() else '_' for ch in str(sku)).lower()


def _read_client_items(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    supplier, items, suppliers = None, [], []
    for row in ws.iter_rows(values_only=True):
        vals = list(row) + [None] * 8
        if not any(v is not None for v in vals):
            continue
        if vals[0] and not any(vals[1:8]):
            if str(vals[0]).strip().lower() != 'supplier':
                supplier = str(vals[0]).strip()
                suppliers.append(supplier)
            continue
        if vals[1] == 'Category' or not vals[2]:
            continue
        cat, sku, desc, note, cost, gcc, uae = vals[1:8]
        items.append({'supplier': supplier, 'category': str(cat).strip(), 'sku': str(sku).strip(),
                      'name': str(desc or sku).strip(), 'note': note, 'cost': cost, 'gcc': gcc, 'uae': uae})
    return items, suppliers


def _num(value):
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None


def build_client_files(items_path, out, cost_currency):
    items, suppliers = _read_client_items(items_path)
    products, uae_rows, gcc_rows, issues = [], [], [], []
    seen = set()
    for item in items:
        if item['sku'] in seen:
            issues.append((item['sku'], item['name'], 'Duplicate SKU - skipped'))
            continue
        seen.add(item['sku'])
        cost, gcc, uae = _num(item['cost']), _num(item['gcc']), _num(item['uae'])
        name = item['name'] + (f" ({item['note']})" if item['note'] else '')
        products.append([
            _xmlid(item['sku']), item['sku'], name, item['category'], 'Goods', 'TRUE', '', 'Units',
            'Power Stop' if item['supplier'] == 'Powerstop' else item['supplier'], '', '', '', '', '', '',
            'stock', cost or '', uae or '', '', item['supplier'], cost or '', cost_currency, '',
        ])
        for price, rows, label in ((uae, uae_rows, 'UAE'), (gcc, gcc_rows, 'GCC')):
            if price:
                rows.append([f'pricelist_{label.lower()}_{_xmlid(item["sku"])[8:]}',
                             f'{"UAE Market" if label == "UAE" else "GCC Market"}', 'Product', item['sku'],
                             'Fixed Price', price])
        missing = [label for label, value in (('cost', cost), ('GCC price', gcc), ('UAE price', uae)) if not value]
        if missing:
            issues.append((item['sku'], item['name'], 'Missing ' + ', '.join(missing)))

    _write_template(os.path.join(out, 'client_products_import.xlsx'), TEMPLATES['products'], products)
    pricelist_columns = [
        ('id', 'Unique external ID (re-import updates the price)', True, None, 'pricelist_uae_16_1184'),
        ('pricelist_id', 'Price list', True, ['UAE Market', 'GCC Market'], 'UAE Market'),
        ('applied_on', 'Apply on', True, ['Product'], 'Product'),
        ('product_tmpl_id', 'Product (internal reference)', True, None, '16-1184'),
        ('compute_price', 'Price type', True, ['Fixed Price'], 'Fixed Price'),
        ('fixed_price', 'Price (AED)', True, None, '54'),
    ]
    _write_template(os.path.join(out, 'client_pricelist_uae_import.xlsx'), pricelist_columns, uae_rows)
    _write_template(os.path.join(out, 'client_pricelist_gcc_import.xlsx'), pricelist_columns, gcc_rows)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data issues'
    ws.append(['SKU', 'Description', 'Issue'])
    for cell in ws[1]:
        cell.font = HEADER
    for issue in issues:
        ws.append(list(issue))
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 40
    wb.save(os.path.join(out, 'client_data_issues.xlsx'))
    vendor_rows = [[f"vendor_{''.join(c if c.isalnum() else '_' for c in s).lower()}", s, 'TRUE', '1', 'overseas',
                    '', cost_currency, '', '', '', ''] for s in suppliers]
    _write_template(os.path.join(out, 'client_vendors_import.xlsx'), TEMPLATES['vendors'], vendor_rows)
    return len(products), len(uae_rows), len(gcc_rows), issues, suppliers


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--items', help='Client "Categories and item details" workbook')
    parser.add_argument('--out', default=os.path.join(os.path.dirname(__file__), '..', 'docs', 'data_load'))
    parser.add_argument('--cost-currency', default='AED', help='Currency of the Cost column in the client file')
    args = parser.parse_args()
    out = os.path.abspath(args.out)
    templates = os.path.join(out, 'templates')
    os.makedirs(templates, exist_ok=True)
    for name, columns in TEMPLATES.items():
        _write_template(os.path.join(templates, f'{name}_template.xlsx'), columns)
    print(f'Templates written to {templates}')
    if args.items:
        n_products, n_uae, n_gcc, issues, suppliers = build_client_files(args.items, out, args.cost_currency)
        print(f'Client files written to {out}: {n_products} products, {n_uae} UAE prices, {n_gcc} GCC prices, '
              f'{len(issues)} rows with data issues; suppliers: {", ".join(suppliers)}')


if __name__ == '__main__':
    main()
